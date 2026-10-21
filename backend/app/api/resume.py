import os
import io
import re
import uuid
import zipfile
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException, Body, Response
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models.resume import ResumeVersion
from app.models.profile import Profile
from app.models.user import User
from app.api.auth import get_current_user
from app.services.storage.local_storage import LocalStorageProvider
from app.services.ai import parser
from app.services.ai.latex_parser import LatexResumeParser, extract_and_parse_latex_zip, strip_latex_formatting
from app.services.ai.parser import (
    validate_resume_text,
    extract_resume,
    merge_resume_to_profile,
    apply_reviewed_profile_data,
    ResumeParseError,
    ResumeTextError,
)
from app.models.job import Job
from app.models.application import Application
from app.services.resume_compiler.candidate_data import load_candidate_library, MODULE_KEYS, MODULE_DISPLAY_NAMES
from app.services.resume_compiler.templates import list_templates, get_template
from app.services.resume_compiler.assembly_engine import generate_assembly_plan
from app.services.resume_compiler.refiner import refine_bullet_truthfully
from app.services.resume_compiler.renderer import render_modular_resume
from app.services.resume_compiler.compiler import compile_latex_securely, extract_and_validate_pdf_readability, create_source_archive
from app.services.resume_compiler.version_service import save_immutable_resume_version, compare_resume_versions

router = APIRouter()
storage = LocalStorageProvider()


class ResumeReviewSaveRequest(BaseModel):
    resume_id: int
    apply_to_master_profile: bool = True
    version_name: Optional[str] = None
    target_role: Optional[str] = None
    notes: Optional[str] = None
    is_master: Optional[bool] = False
    extracted_data: Dict[str, Any]


class VersionCreateRequest(BaseModel):
    name: str
    target_role: Optional[str] = None
    notes: Optional[str] = None
    base_resume_id: Optional[int] = None
    selected_projects: Optional[List[int]] = []
    selected_experiences: Optional[List[int]] = []


class VersionUpdateRequest(BaseModel):
    name: Optional[str] = None
    target_role: Optional[str] = None
    notes: Optional[str] = None
    is_master: Optional[bool] = None
    selected_projects: Optional[List[int]] = None
    selected_experiences: Optional[List[int]] = None


def parse_plain_text_heuristics(text: str) -> Dict[str, Any]:
    contact: Dict[str, Optional[str]] = {
        "full_name": None,
        "email": None,
        "phone": None,
        "github": None,
        "linkedin": None,
        "portfolio": None,
        "current_location": None,
    }

    email_m = re.search(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b', text)
    if email_m:
        contact["email"] = email_m.group(0).strip()

    phone_m = re.search(r'(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{3}\)?[-.\s]?)?\d{3,4}[-.\s]?\d{4}', text)
    if phone_m:
        contact["phone"] = phone_m.group(0).strip()

    gh_m = re.search(r'github\.com/([A-Za-z0-9_.-]+)', text)
    if gh_m:
        contact["github"] = f"https://github.com/{gh_m.group(1)}"

    li_m = re.search(r'linkedin\.com/in/([A-Za-z0-9_.-]+)', text)
    if li_m:
        contact["linkedin"] = f"https://linkedin.com/in/{li_m.group(1)}"

    lines = [line.strip() for line in text.split('\n') if line.strip()]
    if lines:
        first = lines[0]
        if len(first) < 50 and not any(k in first.lower() for k in ["resume", "curriculum", "email", "phone", "profile", "http", "www."]):
            contact["full_name"] = first

    return {
        "profile": contact,
        "education": [],
        "experiences": [],
        "projects": [],
        "skills": [],
        "achievements": [],
        "courses": [],
        "flags": []
    }


@router.post("/upload")
async def upload_resume(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Upload and parse an existing resume.
    Prioritizes LaTeX (.tex) sources and ZIP archives containing LaTeX projects.
    Also supports PDF and DOCX formats.
    Preserves original files in private storage, recognizes macros, flags unparsed items,
    and returns extracted facts for user review before writing to the master candidate profile.
    """
    if not file.filename or not file.filename.strip():
        raise HTTPException(status_code=400, detail="A resume filename is required")

    filename_lower = file.filename.lower()
    allowed_exts = ('.tex', '.zip', '.pdf', '.docx', '.txt', '.md')
    if not any(filename_lower.endswith(ext) for ext in allowed_exts):
        raise HTTPException(
            status_code=400,
            detail="Unsupported file format. Please upload LaTeX (.tex, .zip), PDF (.pdf), or Word (.docx) resumes."
        )

    ext = filename_lower.split('.')[-1]
    content_bytes = await file.read()
    if not content_bytes or not content_bytes.strip():
        raise HTTPException(status_code=400, detail="The resume file is empty")

    text_content = ""
    latex_source = None
    extracted_data: Dict[str, Any] = {}
    flags: List[Dict[str, Any]] = []

    # Validate content and extract text BEFORE writing to storage
    if ext == 'tex':
        if b'\x00' in content_bytes:
            raise HTTPException(status_code=400, detail="File contains binary or null characters")
        try:
            text_content = content_bytes.decode('utf-8')
        except UnicodeDecodeError:
            raise HTTPException(status_code=400, detail="LaTeX source must be UTF-8 encoded text")
        try:
            parser.validate_resume_text(text_content)
        except ResumeTextError as rte:
            raise HTTPException(status_code=400, detail=str(rte))
        latex_source = text_content
        latex_p = LatexResumeParser()
        extracted_data = latex_p.parse_full_latex(text_content)
        flags = extracted_data.get("flags", [])

    elif ext == 'zip':
        if not zipfile.is_zipfile(io.BytesIO(content_bytes)):
            raise HTTPException(status_code=400, detail="Invalid ZIP archive")
        try:
            with zipfile.ZipFile(io.BytesIO(content_bytes), 'r') as zf:
                namelist = zf.namelist()
                if not any(name.lower().endswith('.tex') for name in namelist):
                    raise HTTPException(status_code=400, detail="ZIP archive contains no .tex files")
        except Exception:
            raise HTTPException(status_code=400, detail="Corrupted ZIP archive")

    elif ext in ('txt', 'md'):
        if b'\x00' in content_bytes:
            raise HTTPException(status_code=400, detail="File contains binary or null characters")
        try:
            text_content = content_bytes.decode('utf-8')
        except UnicodeDecodeError:
            raise HTTPException(status_code=400, detail="Text resume must be UTF-8 encoded")
        try:
            parser.validate_resume_text(text_content)
        except ResumeTextError as rte:
            raise HTTPException(status_code=400, detail=str(rte))
        extracted_data = parse_plain_text_heuristics(text_content)

    elif ext == 'pdf':
        try:
            from pdfminer.high_level import extract_text as extract_pdf_text
            text_content = extract_pdf_text(io.BytesIO(content_bytes))
        except Exception:
            raise HTTPException(status_code=400, detail="Could not read PDF file: invalid or corrupted file")
        try:
            parser.validate_resume_text(text_content)
        except ResumeTextError as rte:
            raise HTTPException(status_code=400, detail=str(rte))
        extracted_data = parse_plain_text_heuristics(text_content)

    elif ext == 'docx':
        try:
            from docx import Document as DocxDocument
            document = DocxDocument(io.BytesIO(content_bytes))
            lines = [para.text.strip() for para in document.paragraphs if para.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        lines.append(" | ".join(row_text))
            text_content = '\n'.join(lines)
        except Exception:
            raise HTTPException(status_code=400, detail="Could not read Word document: invalid or corrupted file")
        try:
            parser.validate_resume_text(text_content)
        except ResumeTextError as rte:
            raise HTTPException(status_code=400, detail=str(rte))
        extracted_data = parse_plain_text_heuristics(text_content)

    # Store file in private storage only after successful validation
    unique_filename = f"{current_user.id}_{uuid.uuid4().hex}.{ext}"
    content_type = file.content_type or "application/octet-stream"
    file_path = storage.upload_file(unique_filename, io.BytesIO(content_bytes), content_type)

    if ext == 'zip':
        extract_dir = os.path.join(storage.base_dir, f"{current_user.id}_{uuid.uuid4().hex}_extracted")
        try:
            combined_tex, parsed = extract_and_parse_latex_zip(file_path, extract_dir)
            text_content = combined_tex
            latex_source = combined_tex
            extracted_data = parsed
            flags = parsed.get("flags", [])
        except Exception as ve:
            flags.append({"category": "zip", "field": "archive", "message": str(ve), "severity": "warning"})
            text_content = f"% ZIP archive uploaded: {file.filename}\n"
            extracted_data = {
                "profile": {},
                "education": [],
                "experiences": [],
                "projects": [],
                "skills": [],
                "achievements": [],
                "courses": [],
                "flags": flags
            }

    # Save to Profile and ResumeVersion
    try:
        profile = db.query(Profile).filter(Profile.user_id == current_user.id).first()
        if profile is None:
            profile = Profile(user_id=current_user.id)
            db.add(profile)
        if text_content:
            profile.raw_resume_text = text_content
        if latex_source:
            profile.latex_template = latex_source

        # Check if user has any existing resumes to set first as master
        existing_count = db.query(ResumeVersion).filter(ResumeVersion.user_id == current_user.id).count()
        is_master = (existing_count == 0)

        new_resume = ResumeVersion(
            user_id=current_user.id,
            name=file.filename,
            storage_path=file_path,
            file_type=ext,
            is_master=is_master,
            target_role="General",
            version_number=existing_count + 1,
            latex_source=latex_source,
            parsed_json=extracted_data,
            notes=f"Uploaded as {file.filename}"
        )
        db.add(new_resume)
        db.flush()
        resume_id = new_resume.id
        db.commit()

        return {
            "message": "Resume uploaded successfully. Review and edit the extracted facts before saving to your master profile.",
            "id": resume_id,
            "filename": file.filename,
            "file_type": ext,
            "is_master": is_master,
            "extracted_data": extracted_data,
            "flags": flags,
            "raw_text_preview": text_content[:2000]
        }

    except Exception as exc:
        db.rollback()
        storage.delete_file(file_path)
        raise HTTPException(status_code=500, detail=f"Could not persist resume: {str(exc)}") from exc


@router.post("/review-and-save")
def review_and_save_resume(
    payload: ResumeReviewSaveRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Saves candidate-reviewed facts to the candidate master profile and resume version.
    Ensures ownership and transactional consistency.
    """
    resume = db.query(ResumeVersion).filter(
        ResumeVersion.id == payload.resume_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not resume:
        raise HTTPException(status_code=404, detail="Resume version not found or unauthorized")

    try:
        # Update ResumeVersion metadata
        resume.parsed_json = payload.extracted_data
        if payload.version_name:
            resume.name = payload.version_name
        if payload.target_role:
            resume.target_role = payload.target_role
        if payload.notes:
            resume.notes = payload.notes
        if payload.is_master:
            # Clear is_master on other resumes
            db.query(ResumeVersion).filter(
                ResumeVersion.user_id == current_user.id,
                ResumeVersion.id != resume.id
            ).update({"is_master": False})
            resume.is_master = True

        counts = {}
        if payload.apply_to_master_profile:
            counts = apply_reviewed_profile_data(db, current_user.id, payload.extracted_data)
            # Sync latex_template if available
            if resume.latex_source:
                prof = db.query(Profile).filter(Profile.user_id == current_user.id).first()
                if prof:
                    prof.latex_template = resume.latex_source

        db.commit()
        return {
            "message": "Resume facts reviewed and successfully saved to your master candidate profile.",
            "resume_id": resume.id,
            "counts": counts
        }

    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to save profile facts: {str(exc)}") from exc


@router.get("/{resume_id}/preview")
def get_resume_preview(
    resume_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Retrieves the extracted facts and flags of an existing uploaded resume version for review.
    """
    resume = db.query(ResumeVersion).filter(
        ResumeVersion.id == resume_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found")

    extracted = resume.parsed_json or {}
    flags = extracted.get("flags", []) if isinstance(extracted, dict) else []

    return {
        "resume_id": resume.id,
        "name": resume.name,
        "file_type": resume.file_type,
        "is_master": resume.is_master,
        "target_role": resume.target_role,
        "notes": resume.notes,
        "version_number": resume.version_number,
        "extracted_data": extracted,
        "flags": flags,
        "raw_text_preview": (resume.latex_source or "")[:2000]
    }


@router.post("/autofill")
def autofill_resume(
    payload: Optional[Dict[str, Any]] = Body(default=None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Backwards-compatible autofill endpoint for automated flows and scripts.
    """
    profile = db.query(Profile).filter(Profile.user_id == current_user.id).first()
    try:
        extracted = parser.extract_resume(profile.raw_resume_text if profile else None)
    except ResumeTextError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from None
    except ResumeParseError:
        raise HTTPException(status_code=502, detail="Resume extraction failed. Please try again.") from None

    try:
        counts = parser.merge_resume_to_profile(db, current_user.id, extracted)
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(status_code=500, detail="Could not merge resume into profile. No changes were saved.") from None

    return {"message": "Resume autofill complete. Existing profile data was preserved.", "counts": counts}


# =========================================================================
# MODULAR RESUME CREATION SYSTEM ENDPOINTS
# =========================================================================

class ModularPlanRequest(BaseModel):
    job_id: int
    user_locks: Optional[Dict[str, Any]] = None
    custom_instructions: Optional[str] = None
    target_page_length: Optional[int] = 1
    template_name: Optional[str] = None


class ModularPreviewRequest(BaseModel):
    assembly_plan: Dict[str, Any]
    template_name: Optional[str] = "clean_ats"
    template_options: Optional[Dict[str, Any]] = None
    latex_source: Optional[str] = None


class ModularAssembleSaveRequest(BaseModel):
    job_id: Optional[int] = None
    application_id: Optional[int] = None
    name: Optional[str] = None
    latex_source: Optional[str] = None
    assembly_plan: Dict[str, Any]
    template_name: Optional[str] = "clean_ats"
    template_options: Optional[Dict[str, Any]] = None
    notes: Optional[str] = None
    parent_version_id: Optional[int] = None


class ModularRecompileRequest(BaseModel):
    latex_source: str
    template_options: Optional[Dict[str, Any]] = None


@router.get("/templates")
def get_modular_templates():
    """
    Returns the 5 reusable ATS-compliant resume templates.
    """
    return list_templates()


@router.get("/candidate-library")
def get_candidate_content_library(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Returns all 11 modules of candidate data from the persistent Content Library.
    """
    data = load_candidate_library(current_user.id, db)
    return {
        "candidate_data": data,
        "module_keys": MODULE_KEYS,
        "module_display_names": MODULE_DISPLAY_NAMES,
    }


@router.post("/plan")
def create_modular_resume_plan(
    request: ModularPlanRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Analyzes job requirements and generates a structured resume assembly plan.
    Selects, ranks, and orders candidate modules without altering user locks.
    """
    job = db.query(Job).filter(Job.id == request.job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job opening not found")

    candidate_data = load_candidate_library(current_user.id, db)
    job_dict = {
        "id": job.id,
        "title": job.title,
        "company": job.company,
        "description": job.description or "",
        "required_skills_json": job.required_skills_json or [],
        "tech_stack_json": job.tech_stack_json or [],
    }

    plan = generate_assembly_plan(
        candidate_data=candidate_data,
        job=job_dict,
        user_locks=request.user_locks,
        custom_instructions=request.custom_instructions,
        target_page_length=request.target_page_length or 1,
    )
    if request.template_name:
        plan["template_name"] = request.template_name

    return plan


@router.post("/plan/{job_id}")
def create_modular_resume_plan_by_job(
    job_id: int,
    request: Optional[Dict[str, Any]] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Path-based assembly plan endpoint for compatibility.
    """
    req_obj = ModularPlanRequest(
        job_id=job_id,
        user_locks=(request or {}).get("user_locks"),
        custom_instructions=(request or {}).get("custom_instructions"),
        target_page_length=(request or {}).get("target_page_length", 1),
        template_name=(request or {}).get("template_name")
    )
    return create_modular_resume_plan(req_obj, current_user, db)


@router.post("/preview")
def preview_modular_resume(
    request: ModularPreviewRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Renders and compiles modular resume into a temporary preview PDF with readability validation.
    """
    candidate_data = load_candidate_library(current_user.id, db)
    
    tex_source = request.latex_source
    if not tex_source:
        tex_source = render_modular_resume(
            candidate_data=candidate_data,
            assembly_plan=request.assembly_plan,
            template_name=request.template_name or request.assembly_plan.get("template_name", "clean_ats"),
            template_options=request.template_options or request.assembly_plan.get("template_options"),
        )

    temp_dir = f"data/temp/{current_user.id}"
    os.makedirs(temp_dir, exist_ok=True)
    temp_pdf = os.path.join(temp_dir, f"preview_{uuid.uuid4().hex[:8]}.pdf")
    latest_pdf = os.path.join(temp_dir, "latest_preview.pdf")

    compile_ok, msg = compile_latex_securely(tex_source, temp_pdf)
    readability = extract_and_validate_pdf_readability(temp_pdf)

    try:
        import shutil
        shutil.copy2(temp_pdf, latest_pdf)
    except Exception:
        pass

    import time
    return {
        "success": compile_ok,
        "message": msg,
        "latex_source": tex_source,
        "preview_path": temp_pdf,
        "preview_url": f"/api/v1/resume/preview-temp/pdf?t={int(time.time())}",
        "readability_report": readability,
    }


@router.get("/preview-temp/pdf")
def get_resume_temp_preview_pdf(current_user: User = Depends(get_current_user)):
    """
    Streams the candidate's latest compiled preview PDF.
    """
    pdf_path = f"data/temp/{current_user.id}/latest_preview.pdf"
    if not os.path.exists(pdf_path):
        raise HTTPException(status_code=404, detail="No preview compiled yet")
    return FileResponse(pdf_path, media_type="application/pdf")


@router.post("/assemble-and-save")
def assemble_and_save_resume(
    request: ModularAssembleSaveRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Assembles, compiles, and saves an approved job-specific resume as an immutable version.
    """
    candidate_data = load_candidate_library(current_user.id, db)

    tex_source = request.latex_source
    if not tex_source:
        tex_source = render_modular_resume(
            candidate_data=candidate_data,
            assembly_plan=request.assembly_plan,
            template_name=request.template_name or "clean_ats",
            template_options=request.template_options,
        )

    version_record = save_immutable_resume_version(
        db=db,
        user_id=current_user.id,
        job_id=request.job_id,
        application_id=request.application_id,
        name=request.name or "Tailored Resume",
        latex_source=tex_source,
        assembly_plan=request.assembly_plan,
        template_name=request.template_name or "clean_ats",
        template_options=request.template_options,
        notes=request.notes,
        parent_version_id=request.parent_version_id,
    )

    return {
        "id": version_record.id,
        "version_number": version_record.version_number,
        "name": version_record.name,
        "template_name": version_record.template_name,
        "storage_path": version_record.storage_path,
        "latex_source": version_record.latex_source,
        "readability_report": (version_record.parsed_json or {}).get("readability"),
        "created_at": str(version_record.created_at),
    }


@router.get("/versions/job/{job_id}")
def list_job_resume_versions_alias(
    job_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Lists all saved resume versions for a specific job opening.
    """
    versions = db.query(ResumeVersion).filter(
        ResumeVersion.user_id == current_user.id,
        ResumeVersion.job_id == job_id
    ).order_by(ResumeVersion.version_number.desc()).all()

    return [
        {
            "id": v.id,
            "version_number": v.version_number,
            "name": v.name,
            "template_name": v.template_name,
            "created_at": str(v.created_at),
            "notes": v.notes,
            "has_pdf": bool(v.storage_path and os.path.exists(v.storage_path)),
        }
        for v in versions
    ]


@router.get("/versions/{version_id}/diff")
def diff_resume_version(
    version_id: int,
    compare_with_id: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Computes before-and-after change diffs for a resume version.
    """
    v = db.query(ResumeVersion).filter(
        ResumeVersion.id == version_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not v:
        raise HTTPException(status_code=404, detail="Resume version not found")

    target_compare_id = compare_with_id or v.parent_version_id
    if not target_compare_id:
        # Pick previous version for same job if available
        prev = db.query(ResumeVersion).filter(
            ResumeVersion.user_id == current_user.id,
            ResumeVersion.job_id == v.job_id,
            ResumeVersion.version_number < v.version_number
        ).order_by(ResumeVersion.version_number.desc()).first()
        if prev:
            target_compare_id = prev.id

    if not target_compare_id:
        return {
            "version_id": v.id,
            "version_number": v.version_number,
            "message": "First version for this job. No previous version to compare.",
            "sections": (v.assembly_plan or {}).get("sections", {}),
            "refinements": (v.assembly_plan or {}).get("refinements", {}),
        }

    return compare_resume_versions(db, target_compare_id, version_id)


@router.post("/versions/{version_id}/compile")
def recompile_resume_version(
    version_id: int,
    request: ModularRecompileRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Recompiles an existing version when the user manually modifies the LaTeX source.
    """
    v = db.query(ResumeVersion).filter(
        ResumeVersion.id == version_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not v:
        raise HTTPException(status_code=404, detail="Resume version not found")

    pdf_path = v.storage_path
    compile_ok, msg = compile_latex_securely(request.latex_source, pdf_path)
    readability = extract_and_validate_pdf_readability(pdf_path)

    v.latex_source = request.latex_source
    if request.template_options:
        v.template_options = request.template_options

    parsed = dict(v.parsed_json or {})
    parsed["readability"] = readability
    parsed["compilation_message"] = msg
    v.parsed_json = parsed

    # Update source archive if exists
    zip_path = parsed.get("zip_path")
    if zip_path:
        create_source_archive(request.latex_source, zip_path, {"version": v.version_number})

    db.commit()
    db.refresh(v)

    return {
        "success": compile_ok,
        "message": msg,
        "readability_report": readability,
        "version_id": v.id,
        "version_number": v.version_number,
    }


@router.get("/versions/{version_id}/export-zip")
def export_version_source_zip(
    version_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Downloads source archive (.zip) with resume.tex and metadata.
    """
    v = db.query(ResumeVersion).filter(
        ResumeVersion.id == version_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not v:
        raise HTTPException(status_code=404, detail="Resume version not found")

    zip_path = (v.parsed_json or {}).get("zip_path")
    if not zip_path or not os.path.exists(zip_path):
        storage_dir = f"data/resumes/{current_user.id}"
        os.makedirs(storage_dir, exist_ok=True)
        zip_path = os.path.join(storage_dir, f"source_v{v.version_number}.zip")
        create_source_archive(v.latex_source or "", zip_path, {"version": v.version_number})

    filename = f"{v.name.replace(' ', '_')}_source.zip"
    return FileResponse(zip_path, media_type="application/zip", filename=filename)


@router.get("/versions/{version_id}/download-pdf")
def download_version_pdf(
    version_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Downloads the compiled PDF of a specific resume version.
    """
    v = db.query(ResumeVersion).filter(
        ResumeVersion.id == version_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not v or not v.storage_path or not os.path.exists(v.storage_path):
        raise HTTPException(status_code=404, detail="PDF not found")

    filename = f"{v.name.replace(' ', '_')}.pdf"
    return FileResponse(v.storage_path, media_type="application/pdf", filename=filename)


@router.get("/versions/{version_id}/download-tex")
def download_version_tex(
    version_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Downloads the LaTeX source file (.tex) of a specific resume version.
    """
    v = db.query(ResumeVersion).filter(
        ResumeVersion.id == version_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not v or not v.latex_source:
        raise HTTPException(status_code=404, detail="LaTeX source not found")

    filename = f"{v.name.replace(' ', '_')}.tex"
    return Response(
        content=v.latex_source,
        media_type="text/x-tex",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/versions")
def get_resume_versions(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Returns all resume versions belonging to the candidate, ordered by creation date descending.
    """
    resumes = db.query(ResumeVersion).filter(
        ResumeVersion.user_id == current_user.id
    ).order_by(ResumeVersion.created_at.desc()).all()

    return [
        {
            "id": r.id,
            "name": r.name,
            "file_type": r.file_type,
            "is_master": bool(r.is_master),
            "target_role": r.target_role or "General",
            "version_number": r.version_number or 1,
            "notes": r.notes or "",
            "selected_projects": r.selected_projects or [],
            "selected_experiences": r.selected_experiences or [],
            "has_latex": bool(r.latex_source),
            "created_at": r.created_at,
            "updated_at": r.updated_at,
        }
        for r in resumes
    ]


@router.post("/versions")
def create_resume_version(
    version_req: VersionCreateRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Creates a new targeted resume version (e.g. AI Engineer, Backend Heavy),
    optionally branching from an existing version.
    """
    base_resume = None
    if version_req.base_resume_id:
        base_resume = db.query(ResumeVersion).filter(
            ResumeVersion.id == version_req.base_resume_id,
            ResumeVersion.user_id == current_user.id
        ).first()

    existing_count = db.query(ResumeVersion).filter(ResumeVersion.user_id == current_user.id).count()

    new_v = ResumeVersion(
        user_id=current_user.id,
        name=version_req.name,
        storage_path=base_resume.storage_path if base_resume else "",
        file_type=base_resume.file_type if base_resume else "tex",
        is_master=False,
        target_role=version_req.target_role or "Targeted",
        version_number=existing_count + 1,
        latex_source=base_resume.latex_source if base_resume else None,
        parsed_json=base_resume.parsed_json if base_resume else None,
        notes=version_req.notes or "",
        selected_projects=version_req.selected_projects or [],
        selected_experiences=version_req.selected_experiences or []
    )
    db.add(new_v)
    db.commit()
    db.refresh(new_v)
    return new_v


@router.put("/versions/{resume_id}")
def update_resume_version(
    resume_id: int,
    update_req: VersionUpdateRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Updates resume version metadata, target role, notes, and selected items.
    """
    resume = db.query(ResumeVersion).filter(
        ResumeVersion.id == resume_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not resume:
        raise HTTPException(status_code=404, detail="Resume version not found")

    if update_req.name is not None:
        resume.name = update_req.name
    if update_req.target_role is not None:
        resume.target_role = update_req.target_role
    if update_req.notes is not None:
        resume.notes = update_req.notes
    if update_req.selected_projects is not None:
        resume.selected_projects = update_req.selected_projects
    if update_req.selected_experiences is not None:
        resume.selected_experiences = update_req.selected_experiences
    if update_req.is_master is True:
        db.query(ResumeVersion).filter(
            ResumeVersion.user_id == current_user.id,
            ResumeVersion.id != resume.id
        ).update({"is_master": False})
        resume.is_master = True

    db.commit()
    db.refresh(resume)
    return resume


@router.post("/versions/{resume_id}/set-master")
def set_master_resume(
    resume_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Designates a resume version as the candidate's master resume.
    """
    resume = db.query(ResumeVersion).filter(
        ResumeVersion.id == resume_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not resume:
        raise HTTPException(status_code=404, detail="Resume version not found")

    db.query(ResumeVersion).filter(
        ResumeVersion.user_id == current_user.id,
        ResumeVersion.id != resume.id
    ).update({"is_master": False})

    resume.is_master = True

    # Also update Profile.latex_template if latex source is present
    if resume.latex_source:
        prof = db.query(Profile).filter(Profile.user_id == current_user.id).first()
        if prof:
            prof.latex_template = resume.latex_source

    db.commit()
    return {"message": f"'{resume.name}' is now your master resume.", "is_master": True, "id": resume.id}


@router.delete("/{resume_id}")
def delete_resume(
    resume_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Deletes a resume version and removes stored files safely.
    """
    resume = db.query(ResumeVersion).filter(
        ResumeVersion.id == resume_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found")

    # Check if other versions reference the same file
    other_ref = db.query(ResumeVersion).filter(
        ResumeVersion.id != resume.id,
        ResumeVersion.storage_path == resume.storage_path
    ).first()

    if not other_ref and resume.storage_path:
        storage.delete_file(resume.storage_path)

    db.delete(resume)
    db.commit()
    return {"message": "Deleted successfully"}


@router.get("/download/{resume_id}")
def download_resume(
    resume_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Downloads the private stored resume file.
    Enforces user authentication and ownership.
    """
    resume = db.query(ResumeVersion).filter(
        ResumeVersion.id == resume_id,
        ResumeVersion.user_id == current_user.id
    ).first()

    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found")

    if not os.path.exists(resume.storage_path):
        raise HTTPException(status_code=404, detail="File not found in storage")

    return FileResponse(path=resume.storage_path, filename=resume.name)
